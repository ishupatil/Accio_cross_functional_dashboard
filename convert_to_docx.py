import os
import re
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Error: 'python-docx' library is not installed.")
    print("Please run: pip install python-docx")
    import sys
    sys.exit(1)

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Set default font to Calibri 11pt (standard professional layout)
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    if not os.path.exists(md_path):
        print(f"File {md_path} not found.")
        return
        
    print(f"Converting {md_path} to {docx_path}...")
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_headers = []
    table_rows = []
    
    for line in lines:
        line = line.strip()
        
        # Parse table rows starting with |
        if line.startswith('|'):
            if '---' in line or ':---' in line:
                continue
            parts = [p.strip() for p in line.split('|')[1:-1]]
            if not in_table:
                in_table = True
                table_headers = parts
            else:
                table_rows.append(parts)
            continue
        elif in_table:
            # Table ended, build it in Word doc
            if table_headers:
                table = doc.add_table(rows=1, cols=len(table_headers))
                table.style = 'Light Shading Accent 1'
                hdr_cells = table.rows[0].cells
                for idx, text in enumerate(table_headers):
                    hdr_cells[idx].text = re.sub(r'\*\*(.*?)\*\*', r'\1', text) # strip bold markers in header
                for row_data in table_rows:
                    row_cells = table.add_row().cells
                    for idx, val in enumerate(row_data):
                        if idx < len(row_cells):
                            # clean bold and code backticks
                            val_clean = re.sub(r'\*\*(.*?)\*\*', r'\1', val)
                            val_clean = re.sub(r'`(.*?)`', r'\1', val_clean)
                            row_cells[idx].text = val_clean
            in_table = False
            table_headers = []
            table_rows = []
            doc.add_paragraph() # spacing
            
        if not line:
            continue
            
        # Compile Headers
        if line.startswith('# '):
            h = doc.add_heading(line[2:], level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)
        # Compile Bullet Lists
        elif line.startswith('* ') or line.startswith('- '):
            clean_list = re.sub(r'\*\*(.*?)\*\*', r'\1', line[2:])
            clean_list = re.sub(r'`(.*?)`', r'\1', clean_list)
            doc.add_paragraph(clean_list, style='List Bullet')
        # Compile Numbered Lists
        elif re.match(r'^\d+\.\s', line):
            match = re.match(r'^\d+\.\s(.*)$', line)
            text = match.group(1) if match else line[3:]
            clean_num = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            clean_num = re.sub(r'`(.*?)`', r'\1', clean_num)
            doc.add_paragraph(clean_num, style='List Number')
        # Compile Blockquotes
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run(re.sub(r'\*\*(.*?)\*\*', r'\1', line[2:]))
            run.italic = True
        else:
            # Regular text paragraphs
            # Clean up bold syntax, italic syntax, code backticks, and markdown links
            clean_line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
            clean_line = re.sub(r'\*(.*?)\*', r'\1', clean_line)
            clean_line = re.sub(r'`(.*?)`', r'\1', clean_line)
            clean_line = re.sub(r'\[(.*?)\]\(file:///.*?\)', r'\1', clean_line)
            doc.add_paragraph(clean_line)
            
    doc.save(docx_path)
    print(f"Saved: {docx_path}")

def main():
    docs_to_convert = [
        "project_documents/project_tech_stack_breakdown.md",
        "project_documents/master_tools_and_technologies_inventory.md",
        "project_documents/kpi_mathematical_calculations.md",
        "project_documents/dashboard_master_guide.md",
        "project_documents/multidashboard_integration_architecture.md",
        "project_documents/system_integration_guide.md",
    ]
    
    print("Initializing Markdown to Word (.docx) Converter...\n")
    for md in docs_to_convert:
        if os.path.exists(md):
            docx = md.replace(".md", ".docx")
            convert_md_to_docx(md, docx)
        else:
            print(f"Skipping (not found): {md}")
            
    print("\nConversion complete! Look in the 'project_documents' folder.")

if __name__ == "__main__":
    main()
